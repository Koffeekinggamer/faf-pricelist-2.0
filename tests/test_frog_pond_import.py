"""Frog Pond: wood groups on the collection title, never H\"/Finished as Wood."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook

from backend.frog_pond_import import import_frog_pond_workbook
from wide_import import import_workbook

LIVE = Path(
    "/Volumes/ExternalSSD/FAF-pricebook/viztech-downloads/all-20260717/"
    "Frog_Pond_Furniture/Download_2026_Pricelist_2499.xlsx"
)


def _xlsx(rows: list[list], sheet="1 Weston ") -> bytes:
    book = Workbook()
    tab = book.active
    tab.title = sheet
    for row in rows:
        tab.append(row)
    opts = book.create_sheet("Options")
    opts.append(["Options"])
    opts.append(["Painting add 10%"])
    opts.append(["2-Toned Stains add 5%"])
    opts.append(["Rec. Barnwood Oak add", "30% on 1st Column Price"])
    opts.append(["Clear black Walnut add", "50% on 1st Column Price"])
    book.create_sheet("Markup").append(["Standard Markup", 1.0])
    book.create_sheet("Index").append(["Bedroom Collection Index"])
    buf = BytesIO()
    book.save(buf)
    return buf.getvalue()


WESTON = [
    ["", "", "", "", "", "10% Less for unfinished"],
    [
        "Weston Collection",
        "",
        "",
        "",
        "",
        "Sap Cherry, Oak, Brown Maple, Rustic Cherry",
        "Elm, Cherry, Hickory, QSWO, Hard Maple",
    ],
    ["", "", "", "", "", "Finished", "Finished"],
    ["Item #", "Description", 'H"', 'W"', 'D"'],
    ["100K", "King Bed w/ Regular Footboard", 53, 80, 87, 966, 1256],
    ["105", "3 Drawer Night Stand", 26, 24, 17.5, 478, 621],
]


def test_title_row_wood_groups_bind_finished_columns():
    df = import_frog_pond_workbook(_xlsx(WESTON), filename="frog.xlsx").long_df
    items = df[df["line_kind"].fillna("item") != "addon"]
    assert set(items["part_number"]) == {"100K", "105"}
    woods = set(items["species"].astype(str))
    assert any("Sap Cherry" in w and "Oak" in w for w in woods)
    assert any("Elm" in w and "QSWO" in w for w in woods)
    assert "Rec. Barnwood Oak" in woods
    assert "Clear Black Walnut" in woods or "Clear black Walnut" in woods
    barn = items[
        (items["part_number"] == "100K")
        & (items["species"].astype(str).str.contains("Barnwood", case=False))
        & (items["finish_state"] == "finished")
    ]
    assert not barn.empty
    assert float(barn["base_price"].iloc[0]) == 1255.8
    assert not any(
        w in {'H"', 'W"', 'D"', "Wood / Species", "Finished / Glazed"} or w.startswith("Wood Tier")
        for w in woods
    )
    assert items["species"].fillna("").astype(str).str.strip().ne("").all()
    night = items[(items["part_number"] == "105") & (items["finish_state"] == "finished")]
    group_night = night[~night["species"].astype(str).str.contains("Barnwood|Walnut", case=False)]
    assert sorted(group_night["base_price"]) == [478, 621]
    unf = items[(items["part_number"] == "105") & (items["finish_state"] == "unfinished")]
    group_unf = unf[~unf["species"].astype(str).str.contains("Barnwood|Walnut", case=False)]
    assert sorted(round(p, 1) for p in group_unf["base_price"]) == [430.2, 558.9]


def test_import_workbook_keeps_frog_pond_reader():
    result = import_workbook(
        _xlsx(WESTON),
        vendor="Frog Pond Furniture",
        filename="Download_2026_Pricelist_2499.xlsx",
        preferred_parser="frog_pond_furniture",
    )
    assert result.detected_importer == "frog_pond_furniture"
    keys = {
        str(k).strip().lower()
        for k in result.long_df.get("option_key", []).fillna("")
        if str(k).strip()
    }
    assert any("paint" in k for k in keys)
    assert not any("cherry" in k and "hickory" in k for k in keys)
    assert not any("barnwood" in k for k in keys)
    assert not any("walnut" in k for k in keys)


def test_live_book_has_real_woods_not_dimension_headers():
    if not LIVE.is_file():
        return
    result = import_workbook(
        LIVE.read_bytes(),
        vendor="Frog Pond Furniture",
        filename=LIVE.name,
        preferred_parser="frog_pond_furniture",
    )
    assert result.detected_importer == "frog_pond_furniture"
    items = result.long_df[result.long_df["line_kind"].fillna("item") != "addon"]
    assert "100K" in set(items["part_number"].astype(str))
    assert "105" in set(items["part_number"].astype(str))
    assert "1000K" in set(items["part_number"].astype(str))
    woods = set(items["species"].fillna("").astype(str))
    assert not woods & {'H"', 'W"', 'D"', "H", "W", "D", "Wood / Species", "Finished / Glazed"}
    assert not any(w.startswith("Wood Tier") for w in woods)
    blank = items["species"].fillna("").astype(str).str.strip().eq("")
    assert blank.mean() < 0.01
    weston = items[items["part_number"] == "100K"]
    assert weston["species"].fillna("").astype(str).str.strip().ne("").all()
    keys = {
        str(k).strip().lower()
        for k in result.long_df.get("option_key", []).fillna("")
        if str(k).strip()
    }
    assert any("paint" in k for k in keys)
    assert any("tone" in k for k in keys)
    assert not any("barnwood" in k for k in keys)
    assert not any(k == "rec. barnwood oak" for k in keys)
    live_woods = {str(w) for w in items["species"].fillna("")}
    assert any("Barnwood" in w for w in live_woods)
    tried = {t["sheet"]: t for t in result.sheets_tried}
    assert "Master" not in tried
    assert any(str(s).startswith("1 Weston") for s in tried)
