"""Brookside Heritage Hutches: title-row woods on #16, no right-hand twin."""

from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook

from backend.brookside_import import import_brookside_workbook
from wide_import import import_workbook

LIVE = Path(
    "/Volumes/ExternalSSD/FAF-pricebook/viztech-downloads/all-20260717/"
    "Brookside_Home_Furnishings/Download_2026_Pricelist_3164.xlsx"
)


def _xlsx(hutch_rows: list[list], extra_sheet=None) -> bytes:
    book = Workbook()
    tab = book.active
    tab.title = "Heritage  Hutches"
    for row in hutch_rows:
        tab.append(row)
    if extra_sheet:
        other = book.create_sheet("Frontier")
        for row in extra_sheet:
            other.append(row)
    book.create_sheet("Markup").append(["Standard Markup", 1.0])
    buf = BytesIO()
    book.save(buf)
    return buf.getvalue()


HUTCH = [
    ["Heritage Collection", "", "", "", "", "", "", "", "", "", "Heritage Collection"],
    ["Hutches"],
    [
        "Unit #",
        "Description",
        "Oak ",
        0,
        "S Chy Brown Maple Rus Hick",
        0,
        "Cherry Elm Hickory",
        0,
        "QSWO",
        0,
        "Unit #",
        "Description",
        "Oak ",
        "",
        "S Chy Brown Maple Rus Hick",
        "",
        "Cherry Elm Hickory",
        "",
        "QSWO",
    ],
    [
        "",
        "",
        "Finished prices - For Unfinished Deduct end column % ",
        "",
        "",
        "",
        "",
        "",
        "",
        "Unf",
    ],
    [
        "#16",
        "2 Door Top",
        658.75,
        "",
        724.63,
        "",
        823.44,
        "",
        955.19,
        0.08,
        "#16",
        "2 Door Top",
        658.75,
    ],
    ["#13", "2 Door Base", 809.53, "", 890.48, "", 1011.91, "", 1173.82, 0.08],
    ["MA120", "Arm Chair", 196.3, "", 212.3, "", 224.9, "", 261.3, 30],
]


CHAIR_HDR = [
    ["Chair Pricing"],
    [
        "Unit #",
        "Description",
        "Oak",
        "Sap Chy, Br Maple Rus Hick",
        "Cherry Hickory",
        "QSWO",
        "For Unfinished Deduct:",
        "Unit #",
        "Description",
        "Oak",
    ],
    ["MA120", "Arm Chair", 196.3, 212.3, 224.9, 261.3, 30, "MA120", "Arm Chair", 196.3],
]


def test_heritage_hutches_bind_title_row_woods():
    df = import_brookside_workbook(_xlsx(HUTCH), filename="brookside.xlsx").long_df
    items = df[df["line_kind"].fillna("item") != "addon"]
    sixteen = items[items["part_number"].astype(str) == "#16"]
    assert not sixteen.empty
    assert (sixteen["description"].astype(str) == "2 Door Top").all()
    woods = set(sixteen["species"].astype(str))
    assert "Oak" in woods
    assert any("Sap Cherry" in w and "Brown Maple" in w for w in woods)
    assert any("Cherry" in w and "Hickory" in w for w in woods)
    assert any(w == "QSWO" or "QSWO" in w for w in woods)
    oak_fin = sixteen[
        (sixteen["species"].astype(str) == "Oak") & (sixteen["finish_state"] == "finished")
    ]
    assert len(oak_fin) == 1
    assert float(oak_fin["base_price"].iloc[0]) == 658.75
    oak_unf = sixteen[
        (sixteen["species"].astype(str) == "Oak") & (sixteen["finish_state"] == "unfinished")
    ]
    assert not oak_unf.empty
    assert round(float(oak_unf["base_price"].iloc[0]), 2) == 606.05
    assert not items["part_number"].astype(str).isin(["2 Door Top", "2 Door Base"]).any()


def test_import_workbook_keeps_brookside_reader():
    result = import_workbook(
        _xlsx(HUTCH),
        vendor="Brookside Home Furnishings",
        filename="Download_2026_Pricelist_3164.xlsx",
        preferred_parser="brookside_home_furnishings",
    )
    assert result.detected_importer == "brookside_home_furnishings"
    items = result.long_df[result.long_df["line_kind"].fillna("item") != "addon"]
    assert "#16" in set(items["part_number"].astype(str))
    assert items[items["part_number"].astype(str) == "#16"]["species"].fillna("").ne("").all()


def test_live_heritage_hutches_have_woods():
    if not LIVE.is_file():
        return
    result = import_workbook(
        LIVE.read_bytes(),
        vendor="Brookside Home Furnishings",
        filename=LIVE.name,
        preferred_parser="brookside_home_furnishings",
    )
    items = result.long_df[result.long_df["line_kind"].fillna("item") != "addon"]
    sixteen = items[items["part_number"].astype(str) == "#16"]
    assert not sixteen.empty
    woods = set(sixteen["species"].astype(str))
    assert "Oak" in woods
    assert any("QSWO" in w for w in woods)
    assert sixteen["species"].fillna("").astype(str).str.strip().ne("").all()
    heritage = items[items["collection"].astype(str).str.contains("Heritage", case=False, na=False)]
    heritage_blank = heritage["species"].fillna("").astype(str).str.strip().eq("")
    assert int(heritage_blank.sum()) == 0


def test_chair_unfinished_is_dollar_deduct():
    df = import_brookside_workbook(_xlsx(HUTCH), filename="brookside.xlsx").long_df
    chair = df[df["part_number"].astype(str) == "MA120"]
    oak_unf = chair[
        (chair["species"].astype(str) == "Oak") & (chair["finish_state"] == "unfinished")
    ]
    assert not oak_unf.empty
    assert round(float(oak_unf["base_price"].iloc[0]), 2) == 166.3


def test_sap_chy_header_expands():
    df = import_brookside_workbook(_xlsx(CHAIR_HDR), filename="brookside.xlsx").long_df
    woods = set(df["species"].astype(str))
    assert "Oak" in woods
    assert any("Sap Cherry" in w and "Brown Maple" in w and "Rustic Hickory" in w for w in woods)
    assert not any(re.search(r"(?i)\bchy\b|\bhic\b", w) for w in woods)
