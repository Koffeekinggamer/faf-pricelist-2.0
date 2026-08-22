"""Drop must open every Excel tab before importing or ignoring it."""

from __future__ import annotations

import io

import openpyxl
import pytest

from backend.workbook_sheets import (
    classify_sheet_role,
    excel_engine,
    hidden_product_candidates,
    read_all_sheets,
    read_sheet,
)
from wide_import import import_workbook


def _book(sheets: dict[str, list[list]]) -> bytes:
    wb = openpyxl.Workbook()
    first = True
    for name, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_hidden_collection_tabs_are_reported_as_product_candidates():
    """A hidden collection tab may be real product. Report it, never unhide it."""
    data = _book(
        {
            "Pricelist": [["SKU", "Price"], ["A1", 100]],
            "Bedroom": [["SKU", "Price"], ["B1", 200]],
            "Master": [["internal"]],
            "Markup": [["2.7"]],
            "Index": [["contents"]],
            "Pricelist bk": [["SKU", "Price"], ["A1", 90]],
        }
    )
    wb = openpyxl.load_workbook(io.BytesIO(data))
    for name in ("Bedroom", "Master", "Markup", "Index", "Pricelist bk"):
        wb[name].sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)

    candidates = hidden_product_candidates(buf.getvalue())

    assert candidates == ["Bedroom"]
    # Hidden stays hidden: nothing here is imported.
    assert [view.name for view in read_all_sheets(buf.getvalue())] == ["Pricelist"]


def test_read_all_sheets_classifies_every_tab():
    data = _book(
        {
            "Cover": [["ASHERY OAK WHOLESALE"], ["Fredericksburg OH"]],
            "Markup": [[None, "Markup"], [None, 1]],
            "Options": [["Wood Species Pricing"], ["Hickory", 0.1], ["Options"], ["For painting , Add:", 0.35]],
            "Items": [["Part #", "Description", "Wholesale"], ["P1", "Bench", 220]],
        }
    )
    views = read_all_sheets(data)
    by = {v.name: v.role for v in views}
    assert by["Cover"] == "cover"
    assert by["Markup"] == "markup"
    assert by["Options"] == "options"
    assert by["Items"] == "catalog"
    assert classify_sheet_role("Options&Portal bk", views[2].raw) == "options"


@pytest.mark.parametrize(
    "sheet_name",
    [
        "Add Ons",
        "Customization",
        "Personalization",
        "Features",
        "Product Options",
    ],
)
def test_generic_option_sheet_aliases_use_the_addon_reader(sheet_name):
    data = _book(
        {
            sheet_name: [["Paint", "Add 20%"]],
            "Catalog": [["Part #", "Description", "Oak"], ["T-1", "Table", 100]],
        }
    )

    views = read_all_sheets(data)
    assert {view.name: view.role for view in views}[sheet_name] == "options"

    result = import_workbook(data, vendor="Test Builder", filename="Test.xlsx")
    addons = result.long_df[result.long_df["line_kind"] == "addon"]
    assert "Paint" in set(addons["option_key"])


def test_generic_drop_views_options_instead_of_skipping():
    data = _book(
        {
            "Cover": [["Title"]],
            "Options&Portal": [
                ["Options"],
                ["For painting , Add:", 0.35],
                ["For Locks on Drawers , Add:", 13],
            ],
            "Catalog": [["Part #", "Description", "Oak"], ["T-1", "Table", 100]],
        }
    )
    result = import_workbook(data, vendor="Test Builder", filename="Test_Pricelist.xlsx")
    tried = {s["sheet"]: s for s in result.sheets_tried}
    assert "Cover" in tried
    assert "viewed" in str(tried["Cover"].get("layout") or tried["Cover"].get("note"))
    assert "Options&Portal" in tried
    assert tried["Options&Portal"]["layout"] == "viewed_options"
    assert tried["Options&Portal"]["rows"] >= 1
    addons = result.long_df[result.long_df.get("line_kind", "") == "addon"] if not result.long_df.empty else result.long_df
    if not result.long_df.empty and "line_kind" in result.long_df.columns:
        keys = set(addons["option_key"].astype(str))
        assert "Paint" in keys or "Drawer lock" in keys


def test_read_sheet_uses_excel_engine_not_filename():
    data = _book({"Wholesale": [["Item", "Oak"], ["T1", 100]]})
    raw = read_sheet(data, "Wholesale", header=None)
    assert not raw.empty
    assert excel_engine(data) == "openpyxl"


def test_excel_engine_picks_xlrd_for_ole_and_openpyxl_for_zip():
    assert excel_engine(b"\xd0\xcf\x11\xe0" + b"\x00" * 20) == "xlrd"
    assert excel_engine(b"PK\x03\x04" + b"\x00" * 20) == "openpyxl"
    assert excel_engine(_book({"Sheet1": [["a"]]})) == "openpyxl"


def test_hidden_sheet_is_left_hidden_and_not_imported():
    wb = openpyxl.Workbook()
    vis = wb.active
    vis.title = "Pricelist"
    vis.append(["Part #", "Description", "Oak"])
    vis.append(["V1", "Visible Table", 100])
    vis.append(["V2", "Visible Bench", 140])
    vis.append(["V3", "Visible Chest", 220])
    hid = wb.create_sheet("Master")
    hid.sheet_state = "hidden"
    hid.append(["Part #", "Description", "Oak"])
    hid.append(["H1", "Hidden Dup Table", 100])
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()

    views = read_all_sheets(data)
    assert [v.name for v in views] == ["Pricelist"]
    result = import_workbook(data, vendor="Hide Test", filename="Hide.xlsx")
    parts = set(result.long_df["part_number"].astype(str)) if not result.long_df.empty else set()
    assert "V1" in parts
    assert "H1" not in parts
    tried = {s["sheet"] for s in result.sheets_tried}
    assert "Master" not in tried or "hidden" in str(
        next(s for s in result.sheets_tried if s["sheet"] == "Master").get("note", "")
    ).lower()


def test_hidden_rows_below_visible_catalog_are_not_imported():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pricelist"
    ws.append(["Part #", "Description", "Oak"])
    ws.append(["V1", "Visible Table", 100])
    ws.append(["V2", "Visible Bench", 140])
    ws.append(["V3", "Visible Chest", 220])
    ws.append(["H2", "Dup hidden below", 100])
    ws.row_dimensions[5].hidden = True
    buf = io.BytesIO()
    wb.save(buf)
    result = import_workbook(buf.getvalue(), vendor="Hide Test", filename="HideRows.xlsx")
    parts = set(result.long_df["part_number"].astype(str)) if not result.long_df.empty else set()
    assert "V1" in parts
    assert "H2" not in parts
