"""Kidron's visible Prices tab has retail formulas left and wholesale twins right."""

import io

import openpyxl

from backend.kidron_import import import_kidron_workbook


def _book() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Prices"
    ws.append(
        [
            None,
            "Brown Maple\nOak",
            None,
            "Cherry\nWhite Oak\nHickory\nQSWO",
            None,
            None,
            None,
            None,
            None,
            "Brown Maple\nOak",
            None,
            "QSWO\nClear Maple\nHickory, Cherry",
            None,
            "Rustic\nCherry\nRustic QSWO",
            None,
            "Walnut",
        ]
    )
    ws.append(
        [
            "Solo Bookcases",
            "Fin",
            "Unfin",
            "Fin",
            "Unfin",
            None,
            None,
            None,
            None,
            "Fin",
            "Unfin",
            "Fin",
            "Unfin",
            "Fin",
            "Unfin",
            "Fin",
            "Unfin",
        ]
    )
    ws.append(
        [
            "#S8080 Solo Bookcase",
            2180,
            None,
            2610,
            None,
            None,
            None,
            None,
            None,
            716,
            628,
            943,
            831,
            830,
            729,
            1145,
            1007,
        ]
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _single_band_book(sheet_title: str) -> bytes:
    """Harbor / Occasionals shape: one wholesale band, no retail formulas."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append([None, None, None, None, None, "Oak\nBrown Maple", None, "White Oak"])
    ws.append([None, None, None, None, None, "Fin", "Unfin", "Fin", "Unfin"])
    ws.append(["#601 Mule Dresser", None, None, None, None, 1160, 1040, 1404, 1284])
    ws.append(["All other wood species available."])
    ws.append(["* Sturdy Act Compliant"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_harbor_style_sheet_is_read_even_though_the_tab_is_not_named_prices():
    result = import_kidron_workbook(
        _single_band_book("Sheet1"),
        vendor="Kidron Woodcraft",
        filename="Harbor.xlsx",
    )
    rows = result.long_df

    assert len(rows) == 4
    oak = rows[rows["species"] == "Oak / Brown Maple"]
    assert dict(zip(oak["finish_state"], oak["base_price"])) == {
        "finished": 1160.0,
        "unfinished": 1040.0,
    }
    assert set(rows["part_number"]) == {"#601"}


def test_timberline_puts_the_sku_after_the_description_in_one_cell():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([None, "Brown Maple Oak (Rough Sawn)", None, "Walnut"])
    ws.append([None, "Fin", "Unfin", "Fin", "Unfin"])
    ws.append(["Deluxe Dresser                #1974", 1225, 1092, 1846, 1638])
    buf = io.BytesIO()
    wb.save(buf)

    rows = import_kidron_workbook(
        buf.getvalue(),
        vendor="Kidron Woodcraft",
        filename="Timberline.xlsx",
    ).long_df

    assert len(rows) == 4
    assert set(rows["part_number"]) == {"#1974"}
    assert set(rows["description"]) == {"Deluxe Dresser"}
    walnut = rows[rows["species"] == "Walnut"]
    assert dict(zip(walnut["finish_state"], walnut["base_price"])) == {
        "finished": 1846.0,
        "unfinished": 1638.0,
    }


def test_repeated_species_block_is_not_counted_twice():
    """Entertainment Center repeats one wholesale block; the SKU must not double."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([None, "Brown Maple\nOak", None, None, None, None, None, None, None, "Brown Maple\nOak"])
    ws.append(["***ROUGH SAWN PRICE", "Fin", "Unfin", None, None, None, None, None, None, "Fin", "Unfin"])
    ws.append(["#1962 Timberline TV Stand", 716, 628, None, None, None, None, None, None, 716, 628])
    buf = io.BytesIO()
    wb.save(buf)

    rows = import_kidron_workbook(
        buf.getvalue(),
        vendor="Kidron Woodcraft",
        filename="Timberline_Entertainment_Center.xlsx",
    ).long_df

    assert len(rows) == 2
    assert not any("ROUGH SAWN" in str(value) for value in rows["collection"])


def test_distinct_species_blocks_both_survive():
    """Timberline's right block is a different wood tier, not a repeat."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([None, "Brown Maple Oak (Rough Sawn)", None, None, None, None, None, None, None, "Walnut"])
    ws.append([None, "Fin", "Unfin", None, None, None, None, None, None, "Fin", "Unfin"])
    ws.append(["Deluxe Dresser #1974", 1225, 1092, None, None, None, None, None, None, 1846, 1638])
    buf = io.BytesIO()
    wb.save(buf)

    rows = import_kidron_workbook(
        buf.getvalue(),
        vendor="Kidron Woodcraft",
        filename="Timberline.xlsx",
    ).long_df

    assert len(rows) == 4
    assert set(rows["species"]) == {"Brown Maple Oak (Rough Sawn)", "Walnut"}


def test_grand_isle_finished_configurations_are_not_dropped_for_lacking_skus():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(
        [
            None,
            None,
            None,
            None,
            None,
            "With Cowhide",
            None,
            "Without Cowhide",
            None,
            "With Cowhide",
            None,
            "Without Cowhide",
        ]
    )
    ws.append(
        [
            None,
            None,
            None,
            None,
            None,
            "Fin",
            None,
            "Fin",
            None,
            "Fin",
            None,
            "Fin",
        ]
    )
    ws.append(
        [
            "Queen & Full",
            None,
            None,
            None,
            None,
            4425,
            None,
            4125,
            None,
            4425,
            None,
            4125,
        ]
    )
    ws.append(
        [
            "King",
            None,
            None,
            None,
            None,
            4900,
            None,
            4500,
            None,
            4900,
            None,
            4500,
        ]
    )
    buf = io.BytesIO()
    wb.save(buf)

    rows = import_kidron_workbook(
        buf.getvalue(),
        vendor="Kidron Woodcraft",
        filename="Grand_Isle.xlsx",
    ).long_df

    assert len(rows) == 4
    assert set(rows["collection"]) == {"Grand Isle"}
    assert set(rows["description"]) == {
        "Queen & Full — With Cowhide",
        "Queen & Full — Without Cowhide",
        "King — With Cowhide",
        "King — Without Cowhide",
    }
    assert set(rows["finish_state"]) == {"finished"}


def test_collection_falls_back_to_the_book_name():
    rows = import_kidron_workbook(
        _single_band_book("Sheet1"),
        vendor="Kidron Woodcraft",
        filename="Flint_Ridge.xlsx",
    ).long_df

    assert set(rows["collection"]) == {"Flint Ridge"}


def test_footer_prose_does_not_become_a_collection():
    rows = import_kidron_workbook(
        _single_band_book("Pricing"),
        vendor="Kidron Woodcraft",
        filename="Occasionals.xlsx",
    ).long_df

    collections = {str(value) for value in rows["collection"]}
    assert not any("available" in value.lower() for value in collections)
    assert not any(value.startswith("*") for value in collections)


def test_reader_uses_wholesale_finish_twins_not_left_retail_formulas():
    result = import_kidron_workbook(
        _book(),
        vendor="Kidron Woodcraft",
        filename="Solo Galaxy.xlsx",
    )
    rows = result.long_df

    assert len(rows) == 8
    assert set(rows["finish_state"]) == {"finished", "unfinished"}
    brown = rows[rows["species"] == "Brown Maple / Oak"]
    assert dict(zip(brown["finish_state"], brown["base_price"])) == {
        "finished": 716.0,
        "unfinished": 628.0,
    }
    assert 2180.0 not in set(rows["base_price"])
    assert set(rows["part_number"]) == {"#S8080"}
    assert set(rows["description"]) == {"Solo Bookcase"}
    assert "Rustic Cherry / Rustic QSWO" in set(rows["species"])
