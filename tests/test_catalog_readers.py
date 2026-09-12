"""Each catalog builder has a named reader so next year's Drop reuses it."""

from __future__ import annotations

import io

import openpyxl
import pandas as pd

from backend.builder_parsers import guess_named_parser, identify_reader
from backend.builder_profiles import filename_hints_for, match_profile_vendor
from backend.builder_reader_registry import DEFAULT_READER_REGISTRY
from backend.catalog_readers import (
    CATALOG_SPECS,
    apply_catalog_description_fixes,
    import_catalog_workbook,
    spec_for_vendor,
)
from wide_import import import_workbook


def _book(vendor_on_cover: str) -> bytes:
    wb = openpyxl.Workbook()
    cover = wb.active
    cover.title = "Cover"
    cover.append([vendor_on_cover])
    ws = wb.create_sheet("Catalog")
    ws.append(["Part #", "Description", "Species", "Wholesale"])
    ws.append(["T1", "Side Table", "Oak", 100])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_filename_hints_drop_viztech_download_stems():
    hints = filename_hints_for("Black Horse Furniture", "Download_2026_Pricelist_286628.xlsx")
    assert "black horse furniture" in hints
    assert "download" not in hints
    assert not any(h.startswith("download") for h in hints)


def test_filename_hints_include_catalog_spec_tokens():
    hints = filename_hints_for("Brookside Home Furnishings")
    assert "brookside" in hints
    assert "brookside home furnishings" in hints


def test_catalog_readers_do_not_bind_looks_like_globals():
    import backend.catalog_readers as catalog_readers

    assert not hasattr(catalog_readers, "looks_like_black_horse_furniture")
    entry = DEFAULT_READER_REGISTRY.get("black_horse_furniture")
    assert entry.detect_fn is not None
    assert entry.reader_fn is not None
    assert entry.detector == ""


def test_every_catalog_builder_has_a_specific_reader():
    ids = set(DEFAULT_READER_REGISTRY.specific_ids)
    for spec in CATALOG_SPECS:
        assert spec.parser_id in ids
        assert DEFAULT_READER_REGISTRY.get(spec.parser_id).vendor == spec.vendor


def test_next_year_named_file_resolves_to_the_catalog_reader():
    vendor, parser_id = guess_named_parser("Black Horse Furniture/Download_2027_Pricelist.xlsx")
    assert parser_id == "black_horse_furniture"
    assert vendor == "Black Horse Furniture"


def test_xlsx_cover_name_resolves_when_filename_is_viztech_download():
    data = _book("Black Horse Furniture Wholesale")
    vendor, parser_id = guess_named_parser(
        "Download_2027_Pricelist_99999.xlsx",
        data=data,
    )
    assert parser_id == "black_horse_furniture"
    assert vendor == "Black Horse Furniture"


def test_catalog_reader_tags_and_parses_through_generic_unpivot():
    data = _book("Elite Designs")
    result = import_workbook(
        data,
        filename="Elite Designs 2028 Pricelist.xlsx",
        vendor="Elite Designs",
    )
    assert result.detected_importer == "elite_designs"
    assert not result.long_df.empty


def test_ajs_folder_name_with_html_apostrophe_still_matches():
    vendor, parser_id = guess_named_parser(
        "AJ_039_s_Furniture/Download_2027_Pricelist_Finished.xls"
    )
    assert parser_id == "ajs_furniture"
    assert vendor == "AJ's Furniture"


def test_identify_reader_agrees_with_import_workbook():
    data = _book("Black Horse Furniture Wholesale")
    name = "Download_2027_Pricelist_99999.xlsx"
    vendor, parser_id, source = identify_reader(name, data=data)
    result = import_workbook(data, filename=name)
    assert vendor == "Black Horse Furniture"
    assert parser_id == "black_horse_furniture"
    assert source == "guessed"
    assert result.detected_importer == parser_id
    assert result.parser_source == source


def test_bare_download_filename_does_not_match_a_locked_builder():
    assert match_profile_vendor("Download_2027_Pricelist_111.xlsx") is None


def test_ajs_fabric_tier_is_option_not_part_number():
    parsed = pd.DataFrame(
        [
            {
                "part_number": "Leather",
                "description": "101 CSC",
                "option_key": None,
                "species": "RED OAK & BROWN MAPLE",
                "base_price": 1132.10,
            }
        ]
    )

    fixed = apply_catalog_description_fixes(parsed, "AJ's Furniture")

    assert fixed.loc[0, "part_number"] == "101 CSC"
    assert fixed.loc[0, "description"] == "101 CSC"
    assert fixed.loc[0, "option_key"] == "Leather"


def test_integ_upgrades_are_addons_not_fake_catalog_items():
    parsed = pd.DataFrame(
        [
            {
                "collection": "Upgrades / Options",
                "part_number": "#PS2",
                "description": "#PS2",
                "dimensions": "Power Strip, Two 110v, 3 USB",
                "option_key": None,
                "line_kind": None,
                "base_price": 85.0,
            },
            {
                "collection": "#5200 Wall Unit",
                "part_number": "5200",
                "description": "5200",
                "dimensions": "45” TV Opening",
                "option_key": None,
                "line_kind": "item",
                "base_price": 2466.0,
            },
            {
                "collection": "Upgrades",
                "part_number": "Custom Hardware (Not in stock)",
                "description": "Custom Hardware (Not in stock)",
                "dimensions": "Non-stock hardware is $3 per pull/knob",
                "option_key": None,
                "line_kind": None,
                "base_price": 3.0,
            },
        ]
    )

    fixed = apply_catalog_description_fixes(parsed, "INTEG Wood Products")
    upgrade = fixed.iloc[0]

    assert upgrade["line_kind"] == "addon"
    assert upgrade["option_key"] == "Power Strip, Two 110v, 3 USB"
    assert upgrade["part_number"] == "Power Strip, Two 110v, 3 USB"
    assert fixed.iloc[1]["line_kind"] == "item"
    assert fixed.iloc[2]["option_key"].endswith("Per Pull/Knob")


def test_mirror_lake_unfinished_book_keeps_its_finish_twin():
    parsed = pd.DataFrame(
        [
            {
                "part_number": "ML-10-DMB",
                "description": "10 Drawer Mule Box",
                "finish_state": "finished",
                "base_price": 900.0,
            }
        ]
    )

    fixed = apply_catalog_description_fixes(
        parsed,
        "Mirror Lake Woodworks",
        filename="Download_2026_Unfinished_Pricelist_695524.xls",
    )

    assert fixed.loc[0, "finish_state"] == "unfinished"


def test_integ_section_name_replaces_sheet_name_collection():
    parsed = pd.DataFrame(
        [
            {
                "part_number": "5200",
                "description": "5200",
                "collection": "Wall Units 1",
                "dimensions": '45" TV Opening',
            }
        ]
    )

    fixed = apply_catalog_description_fixes(
        parsed,
        "INTEG Wood Products",
        product_context={("Wall Units 1", "5200"): "#5200 Wall Unit"},
    )

    assert fixed.loc[0, "collection"] == "#5200 Wall Unit"


def test_hermies_style_header_replaces_instruction_collection():
    parsed = pd.DataFrame(
        [
            {
                "part_number": "HTS1100-4260",
                "description": "42 x 60",
                "collection": "Additional shapes available at no extra charge",
            }
        ]
    )

    fixed = apply_catalog_description_fixes(
        parsed,
        "Hermies Table Shop",
        product_context={("", "HTS1100-4260"): "652 Mission Double Pedestal Table"},
    )

    assert fixed.loc[0, "collection"] == "652 Mission Double Pedestal Table"


def _millcraft_suite_book() -> bytes:
    """Visible Pricelist shape: suite banner, selected beds, then casegoods."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pricelist"
    for row in (
        ["", "Description", "Item Number", "Standard Wood", "Premium Wood"],
        ["", "Camden Collection"],
        ["", "PANEL BED"],
        ["", "Queen", "MDC56QN", 1399, 1679],
        ["", "LEATHER UPHOLSTERED BED"],
        ["", "Queen", "MDC52QN", 1399, 1679],
        ["", "Camden Collection"],
        ["", "CASEGOODS"],
        ["", "1 Door Nightstand", "MDC26NS", 499, 599],
    ):
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_millcraft_selected_beds_share_the_suite_collection():
    spec = spec_for_vendor("Millcraft")
    result = import_catalog_workbook(
        spec,
        _millcraft_suite_book(),
        vendor="Millcraft",
        filename="Millcraft.xlsx",
    )
    items = result.long_df
    items = items[items["part_number"].isin(["MDC56QN", "MDC52QN", "MDC26NS"])]
    by_part = {
        str(part): set(group["collection"].astype(str))
        for part, group in items.groupby("part_number")
    }
    assert by_part["MDC26NS"] == {"Camden Collection"}
    assert by_part["MDC56QN"] == {"Camden Collection"}
    assert by_part["MDC52QN"] == {"Camden Collection"}
    descs = {
        str(part): set(group["description"].astype(str))
        for part, group in items.groupby("part_number")
    }
    assert any("Panel Bed" in d for d in descs["MDC56QN"])
    assert any("Leather Upholstered Bed" in d for d in descs["MDC52QN"])


def _genuine_oak_finished_unfinished_book() -> bytes:
    wb = openpyxl.Workbook()
    finished = wb.active
    finished.title = "Finished"
    finished.append(["ITEM #", "DESCRIPTION", "Oak, Br. Maple", "Cherry, Hickory"])
    for sku, name, oak, cherry in (
        ("G0-100", "Nightstand", 200, 240),
        ("G0-101", "Dresser", 400, 480),
        ("G0-102", "Chest", 360, 432),
        ("G0-103", "Armoire", 520, 624),
        ("G0-104", "Bookcase", 280, 336),
        ("G0-105", "Media Cabinet", 310, 372),
    ):
        finished.append([sku, name, oak, cherry])
    unfinished = wb.create_sheet("Unfinished")
    unfinished.append(["ITEM #", "DESCRIPTION", "Oak, Br. Maple", "Cherry, Hickory"])
    for sku, name, oak, cherry in (
        ("G0-100", "Nightstand", 160, 192),
        ("G0-101", "Dresser", 320, 384),
        ("G0-102", "Chest", 288, 346),
        ("G0-103", "Armoire", 416, 499),
        ("G0-104", "Bookcase", 224, 269),
        ("G0-105", "Media Cabinet", 248, 298),
    ):
        unfinished.append([sku, name, oak, cherry])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_genuine_oak_unfinished_tab_is_unfinished_not_a_collection():
    spec = spec_for_vendor("Genuine Oak")
    result = import_catalog_workbook(
        spec,
        _genuine_oak_finished_unfinished_book(),
        vendor="Genuine Oak",
        filename="Download_2026_Pricelist_3481.xlsx",
    )
    items = result.long_df
    if "line_kind" in items.columns:
        items = items[items["line_kind"].fillna("item").astype(str).str.lower() != "addon"]
    finishes = set(items["finish_state"].astype(str).str.lower())
    assert finishes == {"finished", "unfinished"}
    unfinished = items[items["finish_state"].astype(str).str.lower() == "unfinished"]
    assert not unfinished.empty
    collections = {str(value).strip().lower() for value in unfinished["collection"].dropna()}
    assert "unfinished" not in collections
    oak = unfinished[unfinished["species"].astype(str).str.contains("Oak", case=False)]
    assert float(oak["base_price"].iloc[0]) == 160.0
