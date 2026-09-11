"""Windy Acres named reader: ITEM # can sit in column 1; woods stay Wood."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook

from wide_import import import_windy_acres_workbook, import_workbook

LIVE = Path(
    "/Volumes/ExternalSSD/FAF-pricebook/viztech-downloads/all-20260717/"
    "Windy_Acres/Download_2026_Pricelist_2396.xlsx"
)


def _xlsx(rows: list[list], sheet="Bedroom Collection") -> bytes:
    book = Workbook()
    tab = book.active
    tab.title = sheet
    for row in rows:
        tab.append(row)
    cover = book.create_sheet("Instructions")
    cover.append(["INSTRUCTIONS - PLEASE READ CAREFULLY!"])
    markup = book.create_sheet("MarkUp")
    markup.append(["Markup", 1.0])
    buf = BytesIO()
    book.save(buf)
    return buf.getvalue()


BEDROOM_BLOCK = [
    ["", "Options"],
    ["", "For two tone add 35.00 per piece"],
    ["", "Painting add 15%"],
    ["", "Paint & Glaze add 25%"],
    ["", "For hidden jewelry drawer add 100.00"],
    ["", "For platform bed add 200.00"],
    ["", "Headboard only cost is 50% of bed"],
    ["", "Addie Collection - Value Groups"],
    [
        "",
        "Addie",
        "",
        "",
        "",
        "",
        "Br. Maple\nWormy Maple\nR. Cherry",
        "",
        "Elm\nHickory\nCherry",
        "",
        "R. Walnut\nQSWO",
    ],
    [
        "",
        "ITEM #",
        "Description",
        'D"',
        'W"',
        'H"',
        "Finished",
        "Unfinished",
        "Finished",
        "Unfinished",
        "Finished",
        "Unfinished",
    ],
    [
        "1702-NS-AC-F",
        "1702",
        "Night stand 3-Drawer",
        16.25,
        22,
        27.75,
        429,
        377,
        455,
        407,
        526,
        475,
    ],
    ["", "Addie Beds w/ Storage Rails"],
    [
        "",
        "ITEM #",
        "Description",
        "",
        'HB H"',
        'FB H"',
        "Finished",
        "Unfinished",
        "Finished",
        "Unfinished",
        "Finished",
        "Unfinished",
    ],
    [
        "1715-CK-AC-WSR-AW",
        "1715-CK",
        "California King-W/Storage Rails",
        "",
        53,
        25,
        1629,
        1438,
        1782,
        1591,
        2013,
        1822,
    ],
]


def test_item_hash_in_column_one_reads_woods_and_finish_pairs():
    data = _xlsx(BEDROOM_BLOCK)
    result = import_windy_acres_workbook(data, filename="Download_2026_Pricelist_2396.xlsx")
    df = result.long_df
    if "line_kind" in df.columns:
        items = df[df["line_kind"].fillna("item") != "addon"]
    else:
        items = df
    assert not items.empty
    parts = set(items["part_number"].astype(str))
    assert "1702" in parts
    assert "1715-CK" in parts
    night = items[items["part_number"] == "1702"]
    assert set(night["finish_state"]) == {"finished", "unfinished"}
    woods = set(night["species"].astype(str))
    assert any("Maple" in w and "Cherry" in w for w in woods)
    assert any("Walnut" in w or "QSWO" in w for w in woods)
    assert not any(w.startswith("Wood Tier") for w in woods)
    # Follow-on ITEM # block with no wood row still keeps the prior woods.
    beds = items[items["part_number"] == "1715-CK"]
    assert not beds.empty
    assert not any(str(w).startswith("Wood Tier") for w in beds["species"])
    dims = str(beds["dimensions"].iloc[0])
    assert 'HB H"53' in dims
    assert 'FB H"25' in dims


def test_import_workbook_keeps_locked_windy_acres_not_generic():
    data = _xlsx(BEDROOM_BLOCK)
    result = import_workbook(
        data,
        vendor="Windy Acres Furniture",
        filename="Download_2026_Pricelist_2396.xlsx",
        preferred_parser="windy_acres",
    )
    assert result.detected_importer == "windy_acres"
    assert not result.long_df.empty
    keys = {
        str(k).strip().lower()
        for k in result.long_df.get("option_key", []).fillna("")
        if str(k).strip()
    }
    assert any("two" in k and "tone" in k for k in keys)
    assert any("paint" in k for k in keys)
    assert any("jewelry" in k for k in keys)
    assert any("platform" in k for k in keys)
    assert any("headboard" in k for k in keys)


def test_live_book_reads_visible_bedroom_not_hidden_master():
    if not LIVE.is_file():
        return
    data = LIVE.read_bytes()
    result = import_workbook(
        data,
        vendor="Windy Acres Furniture",
        filename=LIVE.name,
        preferred_parser="windy_acres",
    )
    assert result.detected_importer == "windy_acres"
    items = result.long_df[result.long_df["line_kind"].fillna("item") != "addon"]
    assert "1702" in set(items["part_number"].astype(str))
    assert items["species"].fillna("").astype(str).str.strip().ne("").all()
    assert items["species"].astype(str).str.startswith("Wood Tier").mean() < 0.2
    assert not items["species"].astype(str).str.contains(r"R\.\s*/").any()
    tried = {t["sheet"]: t for t in result.sheets_tried}
    assert "Bedroom Collection" in tried
    assert tried["Bedroom Collection"]["rows"] > 0
    assert "Master" not in tried
    keys = {
        str(k).strip().lower()
        for k in result.long_df.get("option_key", []).fillna("")
        if str(k).strip()
    }
    assert keys, "Bedroom Collection Options must become Search Options"
    assert any("two" in k and "tone" in k for k in keys)
    assert any("jewelry" in k for k in keys)
    assert any("platform" in k for k in keys)
    assert any("headboard" in k for k in keys)
