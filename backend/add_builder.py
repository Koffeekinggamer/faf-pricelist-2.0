"""Scaffold the known path for adding builder N (ADR-0011 / audit job 4).

A later SETTLED-expansion job clones this: plan a canonical vendor, write the
profile JSON + synthetic Options fixture, splice the CatalogSpec or
ReaderEntry, and keep the SETTLED row as a placeholder until a shape test
exists. This module does not invent selling factories and never writes
``generic`` under a specific lock.
"""

from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Optional, Sequence

from backend.builder_identity import (
    WATCHED_SHORT_TOKENS,
    identity_claims,
    is_anonymous_download_label,
    is_short_token,
    is_viztech_download_stem,
)
from backend.builder_profiles import filename_hints_for, vendor_slug
from backend.builder_reader_registry import DEFAULT_READER_REGISTRY, ReaderEntry
from backend.catalog_readers import CatalogSpec

STUB_VENDOR = "Stub Workshop"
STUB_PARSER_ID = "stub_workshop"
STUB_OPTION_KEYS = (
    'Size change (up to 10")',
    "Two-tone",
    "Lock",
)
STUB_NEXT_FILE = "Stub_Workshop_2028_Pricelist.xlsx"
STUB_SHEETS = ("Cover", "Options", "Pricelist")
# Clone into tests/test_builder_parser_contract.SETTLED only after a shape
# test exists. None detector/live-book slots keep this out of the contract.
STUB_SETTLED_PLACEHOLDER = (
    STUB_VENDOR,
    STUB_PARSER_ID,
    STUB_NEXT_FILE,
    list(STUB_SHEETS),
    None,
    None,
)

_GENERIC = frozenset({"generic", "pdf"})


class AddBuilderError(ValueError):
    """User-facing refusal while planning or writing a builder scaffold."""


@dataclass(frozen=True)
class BuilderScaffold:
    vendor: str
    slug: str
    parser_id: str
    kind: str
    extra_tokens: tuple[str, ...]
    next_year_filename: str
    sheets: tuple[str, ...]
    option_keys: tuple[str, ...]
    profile: dict[str, Any]
    catalog_spec: Optional[CatalogSpec]
    reader_entry: Optional[ReaderEntry]
    settled_placeholder: tuple
    fixture_relpath: str


@dataclass(frozen=True)
class WrittenScaffold:
    profile_path: Path
    fixture_path: Path
    hook_path: Path
    splice_path: Path


def existing_reader_for(vendor: str) -> Optional[ReaderEntry]:
    key = (vendor or "").strip().lower()
    if not key:
        return None
    for entry in DEFAULT_READER_REGISTRY.entries:
        if entry.specific and entry.vendor.strip().lower() == key:
            return entry
    return None


def parser_id_for(vendor: str, parser_id: str = "") -> str:
    raw = (parser_id or "").strip().lower()
    if raw:
        return raw
    existing = existing_reader_for(vendor)
    if existing:
        return existing.parser_id
    return vendor_slug(vendor).replace("-", "_")


def _refuse_download_vendor(vendor: str) -> None:
    if is_viztech_download_stem(vendor) or is_viztech_download_stem(f"{vendor}.xlsx"):
        raise AddBuilderError(f"vendor cannot be a Viztech Download stem: {vendor}")
    if is_anonymous_download_label(vendor):
        raise AddBuilderError(f"vendor cannot be a Viztech Download stem: {vendor}")


def _canonical_vendor(vendor: str) -> str:
    raw = (vendor or "").strip()
    if not raw:
        raise AddBuilderError("vendor is required")
    _refuse_download_vendor(raw)
    if is_short_token(raw) and raw.casefold() in WATCHED_SHORT_TOKENS:
        raise AddBuilderError(f"vendor cannot be a watched short token: {raw}")
    from backend.standardize import resolve_builder_vendor

    resolved = resolve_builder_vendor(raw) or raw
    resolved = str(resolved).strip()
    _refuse_download_vendor(resolved)
    return resolved


def _validate_parser_id(parser_id: str) -> str:
    pid = (parser_id or "").strip().lower()
    if not pid:
        raise AddBuilderError("parser id is required")
    if pid in _GENERIC or pid in DEFAULT_READER_REGISTRY.generic_ids:
        raise AddBuilderError(f"importer cannot be generic: {pid}")
    return pid


def _validate_extra_tokens(
    vendor: str,
    extra_tokens: Sequence[str] = (),
) -> tuple[str, ...]:
    claimed: dict[str, set[str]] = {}
    for claim in identity_claims():
        if claim.vendor == vendor or not claim.token:
            continue
        claimed.setdefault(claim.token, set()).add(claim.vendor)

    out: list[str] = []
    seen: set[str] = set()
    for raw in extra_tokens:
        token = " ".join(str(raw or "").strip().lower().split())
        if not token or token in seen:
            continue
        if token in WATCHED_SHORT_TOKENS:
            raise AddBuilderError(f"extra token collides with watched short token: {token}")
        owners = claimed.get(token) or set()
        if owners:
            other = ", ".join(sorted(owners))
            raise AddBuilderError(f"extra token {token} already claimed by {other}")
        seen.add(token)
        out.append(token)
    return tuple(out)


def _safe_filename_hints(vendor: str, source_file: str = "") -> list[str]:
    hints: list[str] = []
    for hint in filename_hints_for(vendor, source_file):
        text = str(hint).strip().lower()
        if not text or is_short_token(text) or text in WATCHED_SHORT_TOKENS:
            continue
        if text not in hints:
            hints.append(text)
    return hints


def _profile_dict(
    vendor: str,
    parser_id: str,
    *,
    kind: str,
    source_file: str,
    stub: bool,
) -> dict[str, Any]:
    layouts = ["wide_species"] if kind == "catalog" else [parser_id]
    out: dict[str, Any] = {
        "version": 1,
        "vendor": vendor,
        "parser": {
            "importer": parser_id,
            "source_file": source_file,
            "filename_hints": _safe_filename_hints(vendor, source_file),
            "layouts": layouts,
            "locked": True,
        },
    }
    if stub:
        out["stub"] = True
        out["note"] = (
            "Template only - not a selling factory. Clone via scripts/add_builder.py."
        )
    return out


def plan_builder(
    vendor: str,
    *,
    kind: str = "catalog",
    extra_tokens: Sequence[str] = (),
    parser_id: str = "",
    source_file: str = "",
) -> BuilderScaffold:
    """Plan registry + profile + fixture + SETTLED placeholder for one vendor."""
    kind_key = (kind or "catalog").strip().lower()
    if kind_key not in {"catalog", "shape"}:
        raise AddBuilderError(f"unknown kind: {kind}")
    vend = _canonical_vendor(vendor)
    pid = _validate_parser_id(parser_id_for(vend, parser_id))
    extras = _validate_extra_tokens(vend, extra_tokens)
    existing = existing_reader_for(vend)
    slug = vendor_slug(vend)
    next_file = source_file or (
        STUB_NEXT_FILE if vend == STUB_VENDOR else f"{vend.replace(' ', '_')}_2028_Pricelist.xlsx"
    )
    source = source_file or next_file
    stub = vend == STUB_VENDOR
    profile = _profile_dict(
        vend, pid, kind=kind_key, source_file=source, stub=stub
    )
    catalog_spec: Optional[CatalogSpec] = None
    reader_entry: Optional[ReaderEntry] = None
    if kind_key == "catalog":
        catalog_spec = CatalogSpec(pid, vend, extra_tokens=extras)
    else:
        reader_entry = existing or ReaderEntry(
            pid,
            vend,
            f"backend.{pid}_import:looks_like_{pid}",
            f"backend.{pid}_import:import_{pid}_workbook",
            (pid,),
        )
    if vend == STUB_VENDOR:
        placeholder = STUB_SETTLED_PLACEHOLDER
        option_keys = STUB_OPTION_KEYS
        sheets = STUB_SHEETS
        relpath = "tests/fixtures/stubs/stub-workshop.xlsx"
    else:
        placeholder = (vend, pid, next_file, list(STUB_SHEETS), None, None)
        option_keys = STUB_OPTION_KEYS
        sheets = STUB_SHEETS
        relpath = f"tests/fixtures/stubs/{slug}.xlsx"
    return BuilderScaffold(
        vendor=vend,
        slug=slug,
        parser_id=pid,
        kind=kind_key,
        extra_tokens=extras,
        next_year_filename=next_file,
        sheets=sheets,
        option_keys=option_keys,
        profile=profile,
        catalog_spec=catalog_spec,
        reader_entry=reader_entry,
        settled_placeholder=placeholder,
        fixture_relpath=relpath,
    )


def render_catalog_spec(plan: BuilderScaffold) -> str:
    if plan.catalog_spec is None:
        return "# shape kind — splice a ReaderEntry into DEFAULT_READER_REGISTRY"
    extras = plan.catalog_spec.extra_tokens
    if extras:
        return (
            f'CatalogSpec("{plan.parser_id}", "{plan.vendor}", extra_tokens={extras}),'
        )
    return f'CatalogSpec("{plan.parser_id}", "{plan.vendor}"),'


def render_reader_entry(plan: BuilderScaffold) -> str:
    entry = plan.reader_entry
    if entry is None:
        return "# catalog kind — splice a CatalogSpec into CATALOG_SPECS"
    return (
        "ReaderEntry(\n"
        f'    "{entry.parser_id}",\n'
        f'    "{entry.vendor}",\n'
        f'    "{entry.detector}",\n'
        f'    "{entry.reader}",\n'
        f"    {entry.layouts},\n"
        "),"
    )


def render_settled_placeholder(plan: BuilderScaffold) -> str:
    vendor, importer, next_file, sheets, _detector, _live = plan.settled_placeholder
    return (
        "(\n"
        f'    "{vendor}",\n'
        f'    "{importer}",\n'
        f'    "{next_file}",\n'
        f"    {list(sheets)},\n"
        "    None,  # detector — fill when promoting to SETTLED\n"
        "    None,  # live book — optional extra, not the sole proof\n"
        "),"
    )


def render_splice(plan: BuilderScaffold) -> str:
    lines = [
        f"# add_builder splice for {plan.vendor}",
        f"# slug={plan.slug} parser_id={plan.parser_id} kind={plan.kind}",
        "",
        "# 1. Registry — catalog token lock (backend/catalog_readers.py CATALOG_SPECS)",
        render_catalog_spec(plan),
        "",
        "# 2. Registry — shape-specific (backend/builder_reader_registry.py, before catalog entries)",
        render_reader_entry(plan),
        "",
        "# 3. SETTLED placeholder — do NOT splice into SETTLED until a shape test exists",
        render_settled_placeholder(plan),
        "",
        "# 4. Optional VENDOR_CANON + filename hint (backend/standardize.py) for messy stems",
        f'    "{plan.vendor.lower()}": "{plan.vendor}",',
        "",
        "# Empty Options is a capture miss. Prove addon rows from the fixture before claiming Options.",
        "",
    ]
    return "\n".join(lines)


def render_test_hook(plan: BuilderScaffold) -> str:
    option_repr = ", ".join(repr(key) for key in plan.option_keys)
    return (
        f'"""Scaffold hook for {plan.vendor}. Generated by scripts/add_builder.py.\n'
        "\n"
        "Clone: run add_builder for the real factory, splice the CatalogSpec or\n"
        "ReaderEntry, commit the profile + fixture, prove Options, then promote\n"
        "the SETTLED placeholder only after a shape test exists.\n"
        '"""\n'
        "\n"
        "from pathlib import Path\n"
        "\n"
        f'VENDOR = "{plan.vendor}"\n'
        f'PARSER_ID = "{plan.parser_id}"\n'
        f'NEXT_FILE = "{plan.next_year_filename}"\n'
        f"OPTION_KEYS = ({option_repr},)\n"
        f'FIXTURE = Path(__file__).resolve().parent / "fixtures" / "stubs" / "{plan.slug}.xlsx"\n'
        "\n"
        "# SETTLED placeholder — not in tests/test_builder_parser_contract.SETTLED yet.\n"
        "SETTLED_PLACEHOLDER = (\n"
        "    VENDOR,\n"
        "    PARSER_ID,\n"
        "    NEXT_FILE,\n"
        f"    {list(plan.sheets)},\n"
        "    None,\n"
        "    None,\n"
        ")\n"
        "\n"
        "def test_scaffold_fixture_proves_options():\n"
        "    assert FIXTURE.is_file(), \"run scripts/add_builder.py --write and commit the xlsx\"\n"
        "    assert PARSER_ID != \"generic\"\n"
        "    assert OPTION_KEYS, \"empty Options is a capture miss\"\n"
    )


def write_profile(
    plan: BuilderScaffold,
    path: Path,
    *,
    overwrite: bool = False,
) -> Path:
    if path.exists() and not overwrite:
        raise AddBuilderError(f"profile exists: {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(plan.profile, indent=2) + "\n", encoding="utf-8")
    return path


def build_options_fixture_bytes(vendor: str) -> bytes:
    """Visible Cover + Options + wide Pricelist. Options are priced, never empty."""
    import io

    import openpyxl

    wb = openpyxl.Workbook()
    cover = wb.active
    cover.title = "Cover"
    cover["A1"] = f"{vendor} — synthetic add_builder Options fixture"
    options = wb.create_sheet("Options")
    for row in (
        ["", "PREFERRED DEALER Wholesale Price List"],
        ["", "", "Standard Options *No Upcharge*"],
        ["", "Oak", "Brown Maple"],
        ["", "ADD 10%"],
        ["", "Size Changes", "", "", "ADD ON OPTIONS"],
        [
            "",
            "Any casegood or bed can be customized up to 10''.",
            "",
            "",
            "Hidden Jewelry Tray: $30",
        ],
        ["", "Two Toning", "", "", "Lock: $25"],
        ["", "ADD 20%"],
        ["", "Size changes on any casegood or bed over 10''."],
        ["", "Premium Finish Choices:"],
        ["", "Distressing"],
        ["", "Hand Rubbed Oil"],
    ):
        options.append(row)
    prices = wb.create_sheet("Pricelist")
    for row in (
        ["Description", "Item Number", "Oak", "Brown Maple"],
        ["1 Drw Nightstand", "SW22NS", 410.5, 492.5],
        ["6 Drw Chest", "SW66CH", 880, 1056],
    ):
        prices.append(row)
    for sheet in wb.worksheets:
        sheet.sheet_state = "visible"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def write_fixture(plan: BuilderScaffold, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(build_options_fixture_bytes(plan.vendor))
    return path


def write_scaffold(
    plan: BuilderScaffold,
    *,
    repo_root: Path,
    overwrite: bool = False,
) -> WrittenScaffold:
    """Write profile JSON, Options fixture, test hook, and splice notes."""
    root = Path(repo_root)
    profile_path = root / "config" / "builder_profiles" / f"{plan.slug}.json"
    fixture_path = root / plan.fixture_relpath
    hook_path = root / "tests" / f"test_{plan.slug.replace('-', '_')}_builder_scaffold.py"
    splice_path = root / "config" / "builder_stubs" / f"{plan.slug}.splice.txt"
    write_profile(plan, profile_path, overwrite=overwrite)
    write_fixture(plan, fixture_path)
    hook_path.parent.mkdir(parents=True, exist_ok=True)
    hook_path.write_text(render_test_hook(plan), encoding="utf-8")
    splice_path.parent.mkdir(parents=True, exist_ok=True)
    splice_path.write_text(render_splice(plan), encoding="utf-8")
    return WrittenScaffold(
        profile_path=profile_path,
        fixture_path=fixture_path,
        hook_path=hook_path,
        splice_path=splice_path,
    )


def format_plan(plan: BuilderScaffold) -> str:
    lines = [
        f"vendor: {plan.vendor}",
        f"slug: {plan.slug}",
        f"parser_id: {plan.parser_id}",
        f"kind: {plan.kind}",
        f"next_year_filename: {plan.next_year_filename}",
        f"option_keys: {', '.join(plan.option_keys)}",
        f"fixture: {plan.fixture_relpath}",
        "",
        "CatalogSpec:",
        render_catalog_spec(plan),
        "",
        "ReaderEntry:",
        render_reader_entry(plan),
        "",
        "SETTLED placeholder (do not splice until a shape test exists):",
        render_settled_placeholder(plan),
        "",
    ]
    return "\n".join(lines)
