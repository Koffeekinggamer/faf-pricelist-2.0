"""Criswell Bedroom wholesale books (CWF_*.xlsx / Beds A-F / Beds H-W / Living Rooms).

Four files, one vendor. Each collection sheet (and the bed catalog) is printed
twice side-by-side — left is wholesale, right is the markup twin. We keep the
left pane. Markup=1.0 means the numbers already are wholesale.
"""

from __future__ import annotations

import io
import re
from typing import Any, Optional

import pandas as pd
from openpyxl import load_workbook

from wide_import import WorkbookImportResult, list_excel_sheets, tag_import_result

_SKIP_SHEETS = re.compile(
    r"(?i)^(markup|cover|terms|quick\s*ship|index|backcover|back|standard\s*hdw)$"
)
_OPTIONS_SHEET = re.compile(r"(?i)^options")
_FINISH_SHEET = re.compile(r"(?i)finish\s*prices")
_SKU_RE = re.compile(r"(?i)^CWF[\s#\-]*(\d+[A-Za-z0-9\-]*)")
_SIZE_RE = re.compile(
    r"(?i)^(twin|full(?:/queen)?|queen|king(?:/california(?:\s*king)?)?|"
    r"california\s*king|headboard(?:\s*only)?)\b"
)
_BANNER_SKIP = re.compile(
    r"(?i)preferred\s+5\s*piece|please\s+note|standard\s+with|for\s+glaz|"
    r"prices\s+effective|internet\s+policy|buy\s+5|but\s+5"
)


def _bytes_mention_criswell(data: Optional[bytes]) -> bool:
    if not data:
        return False
    from backend.workbook_sheets import excel_engine

    # BIFF .xls is not a zip. Scanning raw OLE bytes for "CWF" false-positives
    # and then load_workbook raises "File is not a zip file".
    if excel_engine(data) != "openpyxl":
        return False
    needles = (b"Criswell", b"CRISWELL", b"CWF8111", b"CWF #")
    try:
        import zipfile

        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            for name in archive.namelist():
                if not name.endswith(".xml"):
                    continue
                blob = archive.read(name)
                if any(n in blob for n in needles):
                    return True
    except (OSError, ValueError):
        return False
    return False


def looks_like_criswell(
    filename: str = "",
    sheet_names: Optional[list[str]] = None,
    data: Optional[bytes] = None,
) -> bool:
    fn = (filename or "").lower()
    if re.search(
        r"criswell|\bcwf\b|beds\s*[ah]\s*-\s*[fw]|living\s*rooms?\s*price",
        fn,
    ):
        return True
    names = [str(n).strip().lower() for n in (sheet_names or [])]
    if any("beds for every taste" in name for name in names):
        return True
    if any("bloomfield collection" in name for name in names):
        return True
    return _bytes_mention_criswell(data)


def _text(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return str(value).replace("\n", " ").strip()


def _price(value: Any) -> Optional[float]:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if pd.isna(value) or float(value) <= 0:
            return None
        return float(value)
    text = _text(value).replace("$", "").replace(",", "")
    if not text or text in {"-", "—", "#REF!"}:
        return None
    try:
        amount = float(text)
    except ValueError:
        return None
    return amount if amount > 0 else None


def _left_pane(row: list[Any]) -> list[Any]:
    cells = list(row or [])
    if not cells:
        return []
    first = _text(cells[0])
    if first:
        for idx in range(4, len(cells)):
            if _text(cells[idx]) == first:
                return cells[:idx]
    if len(cells) > 12:
        return cells[:8]
    return cells


def _normalize_sku(label: str) -> str:
    match = _SKU_RE.match(label or "")
    if not match:
        return ""
    return "CWF" + match.group(1).upper()


def _sheet_rows(data: bytes, sheet_name: str) -> list[list[Any]]:
    from backend.workbook_sheets import hidden_row_indexes

    skip = hidden_row_indexes(data, sheet_name)
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    try:
        if sheet_name not in wb.sheetnames:
            return []
        rows = []
        for i, row in enumerate(wb[sheet_name].iter_rows(values_only=True)):
            if i in skip:
                continue
            rows.append(list(row))
        return rows
    finally:
        wb.close()


def _wood_labels(header_rows: list[list[Any]]) -> dict[int, str]:
    """Wood names per price column.

    Criswell runs footnotes and drawer-unit callouts down the same columns as
    the wood names, so only recognized species are collected here.
    """
    from backend.standardize import standardize_species, wood_species_strict

    collected: dict[int, list[str]] = {}
    for row in header_rows:
        pane = _left_pane(row)
        for idx in range(2, len(pane), 2):
            label = _text(pane[idx])
            if not label or re.fullmatch(r"(?i)finished|unfinished|add", label):
                continue
            canonical = standardize_species(label)
            if not canonical:
                continue
            woods = wood_species_strict(canonical.split(" / "))
            if not woods:
                continue
            bucket = collected.setdefault(idx, [])
            for wood in woods:
                if wood not in bucket:
                    bucket.append(wood)
    return {idx: " / ".join(labels) for idx, labels in collected.items()}


def _emit_priced(
    *,
    vendor: str,
    collection: str,
    part_number: str,
    description: str,
    pane: list[Any],
    woods: dict[int, str],
    line_kind: str = "item",
    option_key: Optional[str] = None,
) -> list[dict]:
    rows: list[dict] = []
    for idx in range(2, len(pane), 2):
        amount = _price(pane[idx] if idx < len(pane) else None)
        if amount is None:
            continue
        species = woods.get(idx) or None
        rows.append(
            {
                "vendor": vendor,
                "collection": collection,
                "part_number": part_number,
                "description": description,
                "option_key": option_key,
                "species": species,
                "finish_state": "finished",
                "base_price": amount,
                "price_basis": "wholesale",
                "line_kind": line_kind,
            }
        )
    return rows


def _parse_options(rows: list[list[Any]], *, vendor: str) -> list[dict]:
    woods = _wood_labels(rows[:6])
    out: list[dict] = []
    for row in rows:
        pane = _left_pane(row)
        label = _text(pane[0] if pane else "")
        if not label or re.fullmatch(r"(?i)options", label):
            continue
        if _SKU_RE.match(label):
            continue
        priced = _emit_priced(
            vendor=vendor,
            collection="Addons",
            part_number=label,
            description=label,
            pane=pane if _price(pane[1] if len(pane) > 1 else None) is None else [label, None, pane[1]],
            woods=woods or {2: "All woods"},
            line_kind="addon",
            option_key=label,
        )
        if not priced:
            # Price often sits in column B on the options sheet.
            amount = _price(pane[1] if len(pane) > 1 else None)
            if amount is None:
                continue
            if re.search(r"(?i)\b(deduct(?:ion)?|less|credit)\b", label):
                amount = -amount
            out.append(
                {
                    "vendor": vendor,
                    "collection": "Addons",
                    "part_number": label,
                    "description": label,
                    "option_key": label,
                    "species": None,
                    "finish_state": "finished",
                    "base_price": amount,
                    "price_basis": "wholesale",
                    "line_kind": "addon",
                    "notes": "Criswell Options",
                }
            )
        else:
            # One addon row per option — take the first wood's price (flat across woods).
            first = priced[0]
            first["species"] = None
            if re.search(r"(?i)\b(deduct(?:ion)?|less|credit)\b", label):
                first["base_price"] = -abs(float(first["base_price"]))
            first["notes"] = "Criswell Options"
            out.append(first)
    return out


def count_criswell_option_lines(data: bytes) -> int:
    """Count priced Options-tab lines without building addon rows.

    Independent of ``_parse_options`` so the Load gate still fires if that
    pass goes silent (ADR-0011). Finish Prices are not this count.
    """
    count = 0
    for name in list_excel_sheets(data):
        if not _OPTIONS_SHEET.match(str(name).strip()):
            continue
        for row in _sheet_rows(data, name):
            pane = _left_pane(row)
            label = _text(pane[0] if pane else "")
            if not label or re.fullmatch(r"(?i)options", label):
                continue
            if _SKU_RE.match(label) or _BANNER_SKIP.search(label):
                continue
            amount = _price(pane[1] if len(pane) > 1 else None)
            if amount is None:
                for idx in range(2, len(pane), 2):
                    amount = _price(pane[idx])
                    if amount is not None:
                        break
            if amount is not None:
                count += 1
    return count


def _parse_product_sheet(
    rows: list[list[Any]],
    *,
    vendor: str,
    sheet_name: str,
) -> list[dict]:
    collection = re.sub(r"(?i)\s*collection\s*$", "", sheet_name).strip() or sheet_name
    if re.search(r"(?i)beds for every taste", sheet_name):
        collection = "Beds"
    woods: dict[int, str] = {}
    header_buf: list[list[Any]] = []
    current_bed = ""
    current_sku = ""
    out: list[dict] = []
    for row in rows:
        pane = _left_pane(row)
        label = _text(pane[0] if pane else "")
        prices = [
            _price(pane[idx])
            for idx in range(2, len(pane), 2)
            if _price(pane[idx]) is not None
        ]
        if not label and not prices:
            if any(_text(pane[idx]) for idx in range(2, min(len(pane), 12), 2)):
                header_buf.append(row)
                woods = _wood_labels(header_buf) or woods
            continue
        if label and not prices:
            if _BANNER_SKIP.search(label):
                continue
            woods = _wood_labels(header_buf + [row]) or woods
            header_buf.append(row)
            sku = _normalize_sku(label)
            if re.search(r"(?i)collection", label) and not sku:
                collection = re.sub(
                    r"(?i)\s*collection.*$", "", label
                ).strip() or collection
            if sku or re.search(r"(?i)\bbed\b", label):
                current_bed = label
                current_sku = sku or current_sku
            elif len(header_buf) < 8:
                pass
            continue
        if not woods:
            woods = _wood_labels(header_buf)
        sku = _normalize_sku(label)
        if sku and prices:
            desc = _text(pane[1] if len(pane) > 1 else "") or label
            current_bed = desc
            current_sku = sku
            coll = collection
            if collection == "Beds":
                coll = re.sub(r"(?i)^CWF[\s#\-]*\d+\s*", "", label).strip() or collection
            out.extend(
                _emit_priced(
                    vendor=vendor,
                    collection=coll,
                    part_number=sku,
                    description=desc,
                    pane=pane,
                    woods=woods,
                )
            )
            header_buf = []
            continue
        if _SIZE_RE.match(label) and prices:
            desc = f"{current_bed} — {label}" if current_bed else label
            part = f"{current_sku} {label}".strip() if current_sku else label
            coll = (
                re.sub(r"(?i)^CWF[\s#\-]*\d+\s*", "", current_bed).strip() or collection
                if collection == "Beds"
                else collection
            )
            out.extend(
                _emit_priced(
                    vendor=vendor,
                    collection=coll or collection,
                    part_number=part,
                    description=desc,
                    pane=pane,
                    woods=woods,
                )
            )
            continue
        if prices and label and not _BANNER_SKIP.search(label):
            desc = _text(pane[1] if len(pane) > 1 else "") or label
            out.extend(
                _emit_priced(
                    vendor=vendor,
                    collection=collection,
                    part_number=label[:80],
                    description=desc,
                    pane=pane,
                    woods=woods,
                )
            )
    return out


def import_criswell_workbook(
    data: bytes,
    *,
    vendor: str = "",
    default_collection: str = "",
    sheet_filter: Optional[list[str]] = None,
    filename: str = "",
) -> WorkbookImportResult:
    names = list_excel_sheets(data)
    vendor_name = vendor or "Criswell Bedroom"
    frames: list[pd.DataFrame] = []
    tried: list[dict] = []
    targets = sheet_filter if sheet_filter is not None else names

    for name in names:
        if name not in targets:
            tried.append({"sheet": name, "layout": "skip", "rows": 0, "note": "filtered"})
            continue
        if _SKIP_SHEETS.match(str(name).strip()):
            tried.append({"sheet": name, "layout": "viewed", "rows": 0, "note": "viewed"})
            continue
        raw = _sheet_rows(data, name)
        if _OPTIONS_SHEET.match(str(name).strip()):
            addons = _parse_options(raw, vendor=vendor_name)
            if addons:
                frames.append(pd.DataFrame(addons))
            tried.append(
                {
                    "sheet": name,
                    "layout": "criswell_options",
                    "rows": len(addons),
                    "note": f"{len(addons)} option rows",
                }
            )
            continue
        if _FINISH_SHEET.search(str(name)):
            addons = _parse_options(raw, vendor=vendor_name)
            for row in addons:
                row["notes"] = "Criswell Finish Prices"
            if addons:
                frames.append(pd.DataFrame(addons))
            tried.append(
                {
                    "sheet": name,
                    "layout": "viewed_options",
                    "rows": len(addons),
                    "note": f"{len(addons)} finish addons",
                }
            )
            continue
        items = _parse_product_sheet(raw, vendor=vendor_name, sheet_name=name)
        if items:
            frames.append(pd.DataFrame(items))
        tried.append(
            {
                "sheet": name,
                "layout": "criswell",
                "rows": len(items),
                "note": default_collection or "",
            }
        )

    long_df = (
        pd.concat(frames, ignore_index=True, sort=False) if frames else pd.DataFrame()
    )
    return tag_import_result(
        WorkbookImportResult(
            sheets_tried=tried,
            long_df=long_df,
            detected_markup=1.0,
            sheet_names=names,
            notes=(
                f"{filename + ': ' if filename else ''}Criswell Bedroom · "
                f"{0 if long_df.empty else len(long_df)} rows"
            ),
            expected_option_lines=count_criswell_option_lines(data),
        ),
        "criswell",
        source="guessed",
    )
