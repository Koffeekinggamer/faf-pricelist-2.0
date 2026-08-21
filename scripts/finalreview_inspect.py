#!/usr/bin/env python3
"""In-memory /finalreview dump: every builder Excel vs parser vs live catalog."""

from __future__ import annotations

import json
import re
import sqlite3
import sys
import traceback
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from backend.book_options import extract_book_options
from backend.builder_parsers import identify_reader
from backend.builder_profiles import load_builder_profile
from backend.config import DB_PATH
from backend.service import PriceBookService
from backend.workbook_sheets import read_all_sheets
from wide_import import import_workbook

CACHE_ROOTS = [
    Path("/Users/lordjudsonmiller/Documents/viztech-downloads"),
    Path("/Users/lordjudsonmiller/Downloads"),
    Path("/Users/lordjudsonmiller/FAF-pricelist-2.0"),
]
EXCEL_EXTS = {".xlsx", ".xls", ".xlsm"}
OPTIONISH = re.compile(
    r"(?i)\b(option|add[\s\-]?on|upgrade|two[\s\-]?tone|glaze|distress|unfinish|"
    r"size change|add(?:ing)?\s+\d+\s*%|package|custom size|paint and glaze|"
    r"hand[\s\-]?rubbed|premium finish)\b"
)
FOOTNOTEISH = re.compile(
    r"(?i)(see notes?|contact for pricing|call for|tbd|to be determined|"
    r"price on request|see other|see tab|refer to)\b"
)
SKUISH = re.compile(r"\b([A-Z]{2,8}\d{2,}[A-Z0-9\-]{0,8})\b")
SKIP_DIR = {".venv", "node_modules", ".git", "__pycache__"}


def index_excels() -> dict[str, list[Path]]:
    out: dict[str, list[Path]] = defaultdict(list)
    for root in CACHE_ROOTS:
        if not root.exists():
            continue
        for p in root.rglob("*"):
            if not p.is_file() or p.suffix.lower() not in EXCEL_EXTS:
                continue
            if p.name.startswith("~$"):
                continue
            if any(part in SKIP_DIR for part in p.parts):
                continue
            out[p.name.lower()].append(p)
    return out


def resolve(source_file: str, index: dict[str, list[Path]]) -> list[Path]:
    found = []
    for part in re.split(r"\s+\+\s+", source_file or ""):
        part = part.strip()
        if not part:
            continue
        hits = index.get(Path(part).name.lower(), [])
        if hits:
            ranked = sorted(
                hits,
                key=lambda h: (0 if "viztech-downloads" in str(h) else 1, len(str(h))),
            )
            found.append(ranked[0])
    return found


def hidden_sheets_from_zip(path: Path) -> list[str]:
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        return []
    try:
        import zipfile
        from xml.etree import ElementTree as ET

        with zipfile.ZipFile(path) as zf:
            xml = zf.read("xl/workbook.xml")
        root = ET.fromstring(xml)
        ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        out = []
        for sh in root.findall(".//m:sheet", ns):
            state = sh.attrib.get("state", "visible")
            if state != "visible":
                out.append(sh.attrib.get("name") or "")
        return [n for n in out if n]
    except Exception:
        return []


def inspect_xls(path: Path) -> dict:
    info = {
        "path": str(path),
        "name": path.name,
        "suffix": path.suffix.lower(),
        "sheets": [],
        "hidden_sheets": [],
        "merged_ranges": 0,
        "comments": 0,
        "formulas": 0,
        "validations": 0,
        "optionish_hits": [],
        "footnote_hits": [],
        "sku_tokens": set(),
        "errors": [],
    }
    try:
        import xlrd

        book = xlrd.open_workbook(str(path), formatting_info=False)
    except Exception as exc:
        info["errors"].append(f"xlrd: {exc}")
        return _jsonable(info)
    for i, name in enumerate(book.sheet_names()):
        ws = book.sheet_by_index(i)
        hidden = bool(getattr(ws, "visibility", 0))
        if hidden:
            info["hidden_sheets"].append(name)
        nonempty = 0
        last_row = 0
        sample_headers = []
        cap = min(ws.nrows, 1200)
        for r in range(cap):
            texts = []
            for c in range(ws.ncols):
                v = ws.cell_value(r, c)
                if v == "" or v is None:
                    continue
                nonempty += 1
                last_row = r + 1
                s = str(v)
                texts.append(s)
                if isinstance(v, str):
                    for m in SKUISH.findall(s.upper()):
                        info["sku_tokens"].add(m)
                    if OPTIONISH.search(s) and len(info["optionish_hits"]) < 40:
                        clip = re.sub(r"\s+", " ", s)[:160]
                        if clip not in info["optionish_hits"]:
                            info["optionish_hits"].append(clip)
                    if FOOTNOTEISH.search(s) and len(info["footnote_hits"]) < 20:
                        clip = re.sub(r"\s+", " ", s)[:160]
                        if clip not in info["footnote_hits"]:
                            info["footnote_hits"].append(clip)
            if r < 8:
                joined = " | ".join(t[:40] for t in texts[:8] if t)
                if joined:
                    sample_headers.append(joined[:200])
        info["sheets"].append(
            {
                "name": name,
                "hidden": hidden,
                "max_row": ws.nrows,
                "max_col": ws.ncols,
                "last_nonempty_row": last_row,
                "nonempty_cells_scanned": nonempty,
                "merged": 0,
                "formulas": 0,
                "comments": 0,
                "validations": 0,
                "header_sample": sample_headers[:6],
                "scanned_rows_cap": cap,
            }
        )
    return _jsonable(info)


def inspect_xlsx(path: Path) -> dict:
    if path.suffix.lower() == ".xls":
        return inspect_xls(path)
    info = {
        "path": str(path),
        "name": path.name,
        "suffix": path.suffix.lower(),
        "sheets": [],
        "hidden_sheets": [],
        "merged_ranges": 0,
        "comments": 0,
        "formulas": 0,
        "validations": 0,
        "optionish_hits": [],
        "footnote_hits": [],
        "sku_tokens": set(),
        "errors": [],
    }
    try:
        wb = openpyxl.load_workbook(
            path, data_only=False, read_only=True, keep_links=False
        )
    except InvalidFileException as exc:
        info["errors"].append(f"openpyxl: {exc}")
        return _jsonable(info)
    except Exception as exc:
        info["errors"].append(f"openpyxl: {exc}")
        return _jsonable(info)
    try:
        for ws in wb.worksheets:
            hidden = getattr(ws, "sheet_state", "visible") != "visible"
            if hidden:
                info["hidden_sheets"].append(ws.title)
            formulas = 0
            nonempty = 0
            last_row = 0
            sample_headers = []
            cap = 1200
            for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if i > cap:
                    break
                texts = []
                for v in row:
                    if v is None or v == "":
                        continue
                    nonempty += 1
                    last_row = i
                    s = str(v)
                    if isinstance(v, str) and v.startswith("="):
                        formulas += 1
                    texts.append(s)
                    if isinstance(v, str):
                        for m in SKUISH.findall(s.upper()):
                            info["sku_tokens"].add(m)
                    if OPTIONISH.search(s) and len(info["optionish_hits"]) < 40:
                        clip = re.sub(r"\s+", " ", s)[:160]
                        if clip not in info["optionish_hits"]:
                            info["optionish_hits"].append(clip)
                    if FOOTNOTEISH.search(s) and len(info["footnote_hits"]) < 20:
                        clip = re.sub(r"\s+", " ", s)[:160]
                        if clip not in info["footnote_hits"]:
                            info["footnote_hits"].append(clip)
                if i <= 8:
                    joined = " | ".join(t[:40] for t in texts[:8] if t)
                    if joined:
                        sample_headers.append(joined[:200])
            info["formulas"] += formulas
            info["sheets"].append(
                {
                    "name": ws.title,
                    "hidden": hidden,
                    "max_row": ws.max_row,
                    "max_col": ws.max_column,
                    "last_nonempty_row": last_row,
                    "nonempty_cells_scanned": nonempty,
                    "merged": 0,
                    "formulas": formulas,
                    "comments": 0,
                    "validations": 0,
                    "header_sample": sample_headers[:6],
                    "scanned_rows_cap": cap,
                }
            )
    except Exception as exc:
        info["errors"].append(f"scan: {exc}")
    finally:
        wb.close()
    zip_hidden = hidden_sheets_from_zip(path)
    for name in zip_hidden:
        if name not in info["hidden_sheets"]:
            info["hidden_sheets"].append(name)
        for sheet in info["sheets"]:
            if sheet["name"] == name:
                sheet["hidden"] = True
    return _jsonable(info)


def _jsonable(info: dict) -> dict:
    info["sku_tokens"] = sorted(info.get("sku_tokens") or [])
    return info


def parse_files(paths: list[Path], vendor: str) -> dict:
    out = {
        "importer": "",
        "parser_source": "",
        "rows": 0,
        "item_rows": 0,
        "addon_rows": 0,
        "parts": [],
        "option_keys": [],
        "finish_states": [],
        "species": [],
        "sheet_roles": [],
        "book_options": [],
        "errors": [],
        "notes": "",
    }
    frames = []
    for path in paths:
        data = path.read_bytes()
        try:
            vend, parser_id, source = identify_reader(path.name, data=data, vendor=vendor)
            out["importer"] = parser_id or out["importer"]
            out["parser_source"] = source or out["parser_source"]
            result = import_workbook(data, vendor=vendor or vend, filename=path.name)
            df = result.long_df
            out["notes"] = (out["notes"] + " | " + (result.notes or "")).strip(" |")
            if df is not None and not df.empty:
                frames.append(df)
            try:
                views = read_all_sheets(data)
                for v in views:
                    out["sheet_roles"].append(
                        {"file": path.name, "sheet": v.name, "role": v.role}
                    )
            except Exception as exc:
                out["errors"].append(f"roles {path.name}: {exc}")
            try:
                extra = extract_book_options(data, vendor=vendor)
                out["book_options"].extend(
                    sorted({str(r.get("option_key") or "") for r in extra if r.get("option_key")})
                )
            except Exception as exc:
                out["errors"].append(f"book_options {path.name}: {exc}")
        except Exception as exc:
            out["errors"].append(f"parse {path.name}: {exc}")
            out["errors"].append(traceback.format_exc()[-400:])
    if not frames:
        return out
    import pandas as pd

    df = pd.concat(frames, ignore_index=True)
    out["rows"] = int(len(df))
    kind = df["line_kind"].fillna("item").astype(str).str.lower() if "line_kind" in df.columns else "item"
    if not isinstance(kind, str):
        out["item_rows"] = int((kind != "addon").sum())
        out["addon_rows"] = int((kind == "addon").sum())
        items = df[kind != "addon"]
    else:
        out["item_rows"] = int(len(df))
        items = df
    if "part_number" in items.columns:
        out["parts"] = sorted({str(x).strip() for x in items["part_number"].dropna() if str(x).strip()})[:400]
    if "option_key" in df.columns:
        out["option_keys"] = sorted({str(x).strip() for x in df["option_key"].dropna() if str(x).strip()})[:200]
    if "finish_state" in df.columns:
        out["finish_states"] = sorted({str(x).strip().lower() for x in df["finish_state"].dropna() if str(x).strip()})
    if "species" in items.columns:
        out["species"] = sorted({str(x).strip() for x in items["species"].dropna() if str(x).strip()})[:80]
    out["book_options"] = sorted(set(out["book_options"]))
    return out


def catalog_snapshot(svc: PriceBookService, vendor: str, conn: sqlite3.Connection) -> dict:
    items = conn.execute(
        """
        SELECT COUNT(*) FROM pricebook
        WHERE vendor=? AND lower(COALESCE(line_kind,'item'))!='addon'
        """,
        (vendor,),
    ).fetchone()[0]
    addons = conn.execute(
        """
        SELECT COUNT(*) FROM pricebook
        WHERE vendor=? AND lower(COALESCE(line_kind,'item'))='addon'
        """,
        (vendor,),
    ).fetchone()[0]
    unfinished = conn.execute(
        """
        SELECT COUNT(*) FROM pricebook
        WHERE vendor=? AND lower(COALESCE(finish_state,''))='unfinished'
          AND lower(COALESCE(line_kind,'item'))!='addon'
        """,
        (vendor,),
    ).fetchone()[0]
    parts = [
        r[0]
        for r in conn.execute(
            """
            SELECT DISTINCT part_number FROM pricebook
            WHERE vendor=? AND TRIM(COALESCE(part_number,''))!=''
              AND lower(COALESCE(line_kind,'item'))!='addon'
            """,
            (vendor,),
        )
    ]
    finish = [
        r[0]
        for r in conn.execute(
            """
            SELECT DISTINCT lower(trim(finish_state)) FROM pricebook
            WHERE vendor=? AND TRIM(COALESCE(finish_state,''))!=''
            """,
            (vendor,),
        )
    ]
    description_quality = conn.execute(
        """
        SELECT
          SUM(CASE WHEN TRIM(COALESCE(description,''))='' THEN 1 ELSE 0 END),
          SUM(CASE WHEN lower(TRIM(COALESCE(description,''))) IN
            ('nan','nat','none','null','option','options','finished','unfinished')
            THEN 1 ELSE 0 END),
          SUM(CASE
            WHEN lower(TRIM(COALESCE(description,''))) =
                 lower(TRIM(COALESCE(part_number,'')))
             AND (
               LENGTH(TRIM(COALESCE(collection,''))) > 3
               OR LENGTH(TRIM(COALESCE(dimensions,''))) > 3
             )
            THEN 1 ELSE 0 END)
        FROM pricebook
        WHERE vendor=? AND lower(COALESCE(line_kind,'item'))!='addon'
        """,
        (vendor,),
    ).fetchone()
    return {
        "item_rows": items,
        "addon_rows": addons,
        "unfinished_rows": unfinished,
        "parts": parts,
        "option_keys": svc.list_option_keys(vendor),
        "finish_states": finish,
        "collections": svc.repo.list_collections(vendor),
        "woods": svc.repo.list_species(vendor=vendor)[:40],
        "description_quality": {
            "blank": int(description_quality[0] or 0),
            "placeholder": int(description_quality[1] or 0),
            "sku_with_context": int(description_quality[2] or 0),
        },
    }


def score(book: dict, parsed: dict, cat: dict, missing_file: bool) -> tuple[str, list[str], list[str], str, str]:
    issues = []
    missed = []
    if missing_file:
        issues.append("Locked source file not found on disk")
        return "Issues found", issues, "Source file missing — cannot prove capture.", "low"
    hidden = []
    for insp in book.get("inspections") or []:
        hidden.extend(insp.get("hidden_sheets") or [])
        if insp.get("errors"):
            issues.append("Workbook open error: " + "; ".join(insp["errors"][:2]))
        if insp.get("footnote_hits"):
            missed.append("Footnote/contact-for-price text: " + "; ".join(insp["footnote_hits"][:4]))
        if insp.get("comments"):
            missed.append(f"{insp['comments']} Excel comments")
        if insp.get("formulas"):
            missed.append(f"{insp['formulas']} formula cells (parser uses cached/openpyxl values, not a recalc)")
    # Hidden sheets stay hidden — they are often leftover duplicates, not misses.
    option_sheets = [
        r
        for r in parsed.get("sheet_roles") or []
        if str(r.get("role") or "").lower() in {"options", "option", "addons"}
    ]
    skip_sheets = [
        r
        for r in parsed.get("sheet_roles") or []
        if str(r.get("role") or "").lower() in {"skip", "cover", "markup"}
    ]
    if option_sheets and cat.get("addon_rows", 0) == 0 and not cat.get("option_keys"):
        issues.append(
            "Options-classified sheet(s) "
            + ", ".join(r["sheet"] for r in option_sheets[:6])
            + " but Search Options empty"
        )
    book_opts = parsed.get("book_options") or []
    search = [s.lower() for s in (cat.get("option_keys") or [])]
    missing_opts = [k for k in book_opts if k.lower() not in search]
    if missing_opts:
        missed.append("Book-extracted Options not in Search: " + ", ".join(missing_opts[:8]))
    optionish = []
    for insp in book.get("inspections") or []:
        optionish.extend(insp.get("optionish_hits") or [])
    if optionish and not cat.get("option_keys"):
        issues.append("Book has option-like text; Search Options empty")
        missed.append("Option-like cells: " + "; ".join(optionish[:5]))
    parsed_parts = set(parsed.get("parts") or [])
    cat_parts = set(cat.get("parts") or [])
    if parsed_parts and cat_parts:
        only_parse = sorted(parsed_parts - cat_parts)[:12]
        only_cat = sorted(cat_parts - parsed_parts)[:12]
        if only_parse:
            issues.append(f"{len(parsed_parts - cat_parts)} parsed SKUs not in live catalog (sample {only_parse[:6]})")
        if only_cat and len(only_cat) > 20:
            missed.append(f"{len(cat_parts - parsed_parts)} catalog SKUs not in this re-parse (multi-file lock or standardize drop)")
    if "unfinished" in (parsed.get("finish_states") or []) and "unfinished" not in (cat.get("finish_states") or []):
        issues.append("Parse sees unfinished; catalog has no unfinished rows")
    excel_unf = any(
        re.search(r"(?i)unfinish", " ".join(insp.get("optionish_hits") or []) + " " + " ".join(s.get("name", "") for s in insp.get("sheets") or []))
        for insp in book.get("inspections") or []
    )
    if excel_unf and cat.get("unfinished_rows", 0) == 0 and "Unfinished" not in (cat.get("option_keys") or []):
        missed.append("Unfinished mentioned in book; catalog has 0 unfinished rows and no Unfinished Option")
    if parsed.get("errors"):
        issues.append("Parse errors: " + "; ".join(parsed["errors"][:2]))
    if not cat.get("item_rows"):
        issues.append("Zero sellable rows in live catalog")
    description_quality = cat.get("description_quality") or {}
    if description_quality.get("blank"):
        issues.append(
            f"{description_quality['blank']} item descriptions are blank"
        )
    if description_quality.get("placeholder"):
        issues.append(
            f"{description_quality['placeholder']} item descriptions are parser placeholders"
        )
    if description_quality.get("sku_with_context"):
        missed.append(
            f"{description_quality['sku_with_context']} SKU-only descriptions have "
            "visible collection/dimension context"
        )
    status = "Pass" if not issues and not missed else "Issues found"
    conf = "high" if status == "Pass" else ("medium" if cat.get("item_rows", 0) > 50 else "low")
    if hidden or missing_opts:
        conf = "medium"
    summary = (
        f"{len(book.get('inspections') or [])} file(s), "
        f"{sum(len(i.get('sheets') or []) for i in book.get('inspections') or [])} sheets, "
        f"parser={parsed.get('importer') or '?'}, "
        f"catalog items={cat.get('item_rows', 0)} addons={cat.get('addon_rows', 0)}, "
        f"Search Options={len(cat.get('option_keys') or [])}"
    )
    return status, issues, missed, summary, conf


def main() -> None:
    index = index_excels()
    svc = PriceBookService()
    svc.ensure_ready()
    conn = sqlite3.connect(DB_PATH)
    db_vendors = [
        r[0]
        for r in conn.execute(
            "SELECT DISTINCT vendor FROM pricebook WHERE vendor!='' ORDER BY vendor"
        )
    ]
    profile_dir = Path("/Users/lordjudsonmiller/FAF-pricelist-2.0/config/builder_profiles")
    vendors = sorted(set(db_vendors) | {load_builder_profile(p.stem.replace("-", " ")) and None or p.stem for p in profile_dir.glob("*.json")})
    # Prefer canonical names from DB / profiles
    profile_vendors = []
    for p in sorted(profile_dir.glob("*.json")):
        data = json.loads(p.read_text())
        profile_vendors.append(data.get("vendor") or p.stem)
    all_vendors = sorted(set(db_vendors) | set(profile_vendors), key=str.lower)

    used_names = set()
    reviews = []
    for vendor in all_vendors:
        prof = load_builder_profile(vendor) or {}
        parser = prof.get("parser") or {}
        source = parser.get("source_file") or ""
        importer_lock = parser.get("importer") or ""
        paths = resolve(source, index)
        for p in paths:
            used_names.add(p.name.lower())
        print(f"START {vendor} source={source!r} files={len(paths)}", flush=True)
        inspections = [inspect_xlsx(p) for p in paths]
        print(f"  inspected {vendor}", flush=True)
        parsed = parse_files(paths, vendor) if paths else {"importer": importer_lock, "errors": ["no source file"], "parts": [], "book_options": [], "sheet_roles": []}
        print(f"  parsed {vendor} importer={parsed.get('importer')}", flush=True)
        cat = catalog_snapshot(svc, vendor, conn)
        status, issues, missed_items, summary, conf = score(
            {"inspections": inspections}, parsed, cat, missing_file=not paths
        )
        sheet_names = []
        for insp in inspections:
            for s in insp.get("sheets") or []:
                sheet_names.append((" [hidden]" if s.get("hidden") else "") + s["name"])
        reviews.append(
            {
                "vendor": vendor,
                "file": source,
                "paths": [str(p) for p in paths],
                "locked_importer": importer_lock,
                "detected_importer": parsed.get("importer"),
                "parser_source": parsed.get("parser_source"),
                "status": status,
                "confidence": conf,
                "structure": summary,
                "sheets": sheet_names,
                "hidden_sheets": [h for insp in inspections for h in (insp.get("hidden_sheets") or [])],
                "merged_ranges": sum(i.get("merged_ranges") or 0 for i in inspections),
                "formulas": sum(i.get("formulas") or 0 for i in inspections),
                "comments": sum(i.get("comments") or 0 for i in inspections),
                "validations": sum(i.get("validations") or 0 for i in inspections),
                "optionish": [h for i in inspections for h in (i.get("optionish_hits") or [])][:12],
                "footnotes": [h for i in inspections for h in (i.get("footnote_hits") or [])][:8],
                "book_options": parsed.get("book_options") or [],
                "search_options": cat.get("option_keys") or [],
                "parse_items": parsed.get("item_rows") or 0,
                "parse_addons": parsed.get("addon_rows") or 0,
                "catalog_items": cat.get("item_rows") or 0,
                "catalog_addons": cat.get("addon_rows") or 0,
                "unfinished": cat.get("unfinished_rows") or 0,
                "finish_states": cat.get("finish_states") or [],
                "collections_n": len(cat.get("collections") or []),
                "issues": issues,
                "missed": missed_items,
                "errors": (parsed.get("errors") or [])[:4],
                "sheet_roles": parsed.get("sheet_roles") or [],
            }
        )
        print(f"{status:12} {vendor:32} files={len(paths)} cat={cat.get('item_rows')} opts={len(cat.get('option_keys') or [])}", flush=True)

    orphan_files = []
    for name, paths in sorted(index.items()):
        if name in used_names:
            continue
        p = paths[0]
        if "FAF-pricelist-2.0" in str(p) and ("tests" in p.parts or ".venv" in p.parts):
            continue
        if "viztech-downloads" not in str(p) and "Downloads" not in str(p):
            continue
        orphan_files.append(str(p))

    report = {
        "vendors_in_catalog": db_vendors,
        "reviews": reviews,
        "orphan_excels": orphan_files[:80],
        "indexed_excel_count": sum(len(v) for v in index.values()),
        "db_path": str(DB_PATH),
    }
    out = Path("/tmp/finalreview.json")
    out.write_text(json.dumps(report, indent=2))
    print("WROTE", out, "reviews", len(reviews), "orphans", len(orphan_files))


if __name__ == "__main__":
    main()
